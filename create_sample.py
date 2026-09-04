"""
Create a sample Word document for testing the index generator.
"""

from docx import Document

# Create a new Document
doc = Document()

# Add title
doc.add_heading('A Brief History of Ancient Civilizations', 0)

# Add several paragraphs with proper nouns
paragraphs = [
    """Ancient Rome was one of the most powerful empires in history. Founded according to 
    legend by Romulus in 753 BCE, Rome grew from a small settlement on the Tiber River into 
    a vast empire that controlled much of Europe, North Africa, and the Middle East.""",
    
    """Julius Caesar, one of Rome's greatest generals and statesmen, played a crucial role in 
    the events that led to the demise of the Roman Republic and the rise of the Roman Empire. 
    His campaigns in Gaul and his crossing of the Rubicon River in 49 BCE changed the course 
    of history.""",
    
    """In Egypt, the ancient city of Alexandria became a center of learning and culture. Founded 
    by Alexander the Great in 331 BCE, Alexandria was home to the famous Library of Alexandria, 
    one of the largest and most significant libraries of the ancient world.""",
    
    """Cleopatra VII, the last active ruler of the Ptolemaic Kingdom of Egypt, was known for 
    her intelligence and political acumen. Her relationships with Julius Caesar and Mark Antony 
    had significant impacts on the politics of Rome.""",
    
    """Ancient Greece gave birth to democracy in the city-state of Athens. Philosophers like 
    Socrates, Plato, and Aristotle laid the foundations of Western philosophy. The city of 
    Sparta, Athens' great rival, was known for its powerful military culture.""",
    
    """The Peloponnesian War between Athens and Sparta lasted from 431 to 404 BCE and ultimately 
    led to the decline of Athenian power. The historian Thucydides provided a detailed account 
    of this conflict in his work.""",
    
    """In ancient Mesopotamia, the city of Babylon reached its height under King Nebuchadnezzar II, 
    who built the famous Hanging Gardens of Babylon, one of the Seven Wonders of the Ancient World. 
    The Code of Hammurabi, one of the oldest deciphered writings of significant length, originated 
    from this region.""",
    
    """The Persian Empire, founded by Cyrus the Great, stretched from the Mediterranean Sea to 
    India. The capital cities of Persepolis and Susa were centers of administration and culture. 
    King Darius I expanded the empire to its greatest extent.""",
    
    """Ancient China saw the rise of great dynasties. The Qin Dynasty, though short-lived, unified 
    China under Emperor Qin Shi Huang, who built the Great Wall of China and the famous Terracotta 
    Army in Xi'an.""",
    
    """The Han Dynasty that followed established the Silk Road, connecting China with Rome and 
    facilitating trade between East and West. Cities like Chang'an (modern-day Xi'an) became 
    major trading centers.""",
    
    """In the Americas, the Maya civilization flourished in present-day Mexico and Central America. 
    Cities like Tikal and Chichen Itza featured impressive pyramids and temples. The Maya developed 
    a sophisticated writing system and made significant advances in astronomy.""",
    
    """The Inca Empire in South America had its capital at Cuzco in modern-day Peru. The emperor 
    Pachacuti expanded the empire significantly. The mountain citadel of Machu Picchu stands as 
    a testament to Inca engineering and architecture.""",
]

# Add all paragraphs
for para_text in paragraphs:
    doc.add_paragraph(para_text)
    doc.add_paragraph('')  # Add spacing

# Save the document
output_file = 'sample_document.docx'
doc.save(output_file)

print(f"✅ Sample document created: {output_file}")
print(f"\nYou can now run:")
print(f"  python3 index_generator.py {output_file}")
